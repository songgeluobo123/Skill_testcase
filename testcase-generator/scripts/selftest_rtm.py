#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""selftest_rtm.py — 双向追溯（RTM）与预审第 10 项的自测。

覆盖：
  1. build_rtm 状态判定：全覆盖 / 部分覆盖 / 未覆盖 / 孤儿 / 疑似映射 / exploratory 豁免
  2. append_rtm_sheet 写入 Excel 第二个 sheet，表头正确
  3. prescreen.py 第 10 项：孤儿默认阻断；--orphan-max 放宽后仅告警
"""
import json
import os
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import coverage_audit as ca

RULES = [
    {"id": "R-001", "module": "登录", "text": "正确密码登录成功", "priority": "P0"},
    {"id": "R-002", "module": "登录", "text": "密码错误提示", "priority": "P1"},
    {"id": "R-003", "module": "账号", "text": "连续错误锁定账号", "priority": "P0"},
]

CASES = [
    {"id": "TC-001", "module": "登录", "title": "正确密码登录", "steps": "输入正确密码点击登录",
     "expected": "应登录成功跳转首页", "test_type": "功能测试", "coverage_rule": "R-001"},
    {"id": "TC-002", "module": "登录", "title": "密码错误登录", "steps": "输入错误密码点击登录",
     "expected": "应提示密码错误", "test_type": "异常测试", "coverage_rule": "R-001"},
    {"id": "TC-003", "module": "登录", "title": "密码错误提示", "steps": "查看密码错误提示",
     "expected": "应显示密码错误文案", "test_type": "功能测试", "coverage_rule": "R-002"},
    {"id": "TC-004", "module": "报表", "title": "导出报表", "steps": "点击导出",
     "expected": "应生成报表文件", "test_type": "功能测试"},
    {"id": "TC-005", "module": "登录", "title": "页面元素展示", "steps": "查看页面元素",
     "expected": "应显示输入框", "test_type": "功能测试"},
    {"id": "TC-006", "module": "探索", "title": "自由探索", "steps": "随意操作",
     "expected": "应无崩溃", "test_type": "功能测试", "exploratory": True},
]

fail = 0


def check(name, cond, detail=""):
    global fail
    print(("  OK   " if cond else "  FAIL ") + name + (f" -> {detail}" if detail and not cond else ""))
    if not cond:
        fail += 1


print("=== 1) build_rtm 状态判定 ===")
rtm = ca.build_rtm(CASES, RULES)
s = rtm["summary"]
print("summary:", s)
check("全覆盖=1", s["full"] == 1, str(s["full"]))
check("部分覆盖=1", s["partial"] == 1, str(s["partial"]))
check("未覆盖=1", s["uncovered"] == 1, str(s["uncovered"]))
check("孤儿=1", s["orphans"] == 1, str(s["orphans"]))
check("疑似映射=1", s["weak_only"] == 1, str(s["weak_only"]))
status = {r["id"]: r["status"] for r in rtm["rules"]}
print("status:", status)
check("R-001 全覆盖", status.get("R-001") == "全覆盖", str(status))
check("R-002 部分覆盖", status.get("R-002") == "部分覆盖", str(status))
check("R-003 未覆盖", status.get("R-003") == "未覆盖", str(status))
check("孤儿为 TC-004", rtm["orphans"][0]["id"] == "TC-004", str(rtm["orphans"]))
check("疑似为 TC-005", rtm["weak_only"][0]["id"] == "TC-005", str(rtm["weak_only"]))
check("TC-006 exploratory 豁免", all(o["id"] != "TC-006" for o in rtm["orphans"]))

print("\n=== 2) Excel 第二 sheet ===")
try:
    from openpyxl import Workbook, load_workbook
    wb = Workbook()
    wb.active.title = "功能测试用例"
    wb.active.append(["用例编号"])
    ca.append_rtm_sheet(wb, rtm)
    check("sheet 名为「需求追溯」", wb.sheetnames[1] == "需求追溯", str(wb.sheetnames))
    tmp = os.path.join(tempfile.gettempdir(), "rtm_selftest.xlsx")
    wb.save(tmp)
    wb2 = load_workbook(tmp)
    ws = wb2["需求追溯"]
    check("第二 sheet 位置正确", wb2.sheetnames[1] == "需求追溯", str(wb2.sheetnames))
    check("表头正确", ws["A1"].value == "需求ID" and ws["E1"].value == "覆盖状态", str(ws["A1"].value))
    check("含孤儿区块", any("孤儿用例" in str(ws.cell(r, 1).value or "") for r in range(1, ws.max_row + 1)))
    wb2.close()
    os.remove(tmp)
except Exception as e:
    check("Excel 第二 sheet", False, repr(e))

print("\n=== 3) prescreen 第 10 项 ===")
with tempfile.TemporaryDirectory() as td:
    cases_path = os.path.join(td, "cases.json")
    rules_path = os.path.join(td, "rules.json")
    with open(cases_path, "w", encoding="utf-8") as f:
        json.dump([
            {"id": "TC-A", "module": "登录", "title": "登录成功", "priority": "P1",
             "test_type": "功能测试", "steps": "输入正确密码", "expected": "应跳转首页",
             "design_method": ["等价类划分"], "coverage_rule": "R-001"},
            {"id": "TC-B", "module": "报表", "title": "导出报表", "priority": "P2",
             "test_type": "功能测试", "steps": "点击导出", "expected": "应生成文件",
             "design_method": ["场景法"]},
        ], f, ensure_ascii=False)
    with open(rules_path, "w", encoding="utf-8") as f:
        json.dump(RULES[:1], f, ensure_ascii=False)   # 只用 1 条规则，保证覆盖率检查通过

    base = [sys.executable, os.path.join(HERE, "prescreen.py"), cases_path,
            "--requirement-rules", rules_path,
            "--p0-min", "0", "--p0-max", "1", "--p2-min", "0", "--p2-max", "1",
            "--scene-min-cases", "99"]
    r1 = subprocess.run(base, capture_output=True, text=True, encoding="utf-8", errors="replace")
    print("  默认阈值 exit=", r1.returncode)
    check("孤儿默认阻断（exit=1）", r1.returncode == 1, str(r1.returncode))
    check("输出含「孤儿用例」", "孤儿用例" in (r1.stdout + r1.stderr), (r1.stdout + r1.stderr)[-300:])

    r2 = subprocess.run(base + ["--orphan-max", "0.6"], capture_output=True, text=True,
                        encoding="utf-8", errors="replace")
    print("  放宽阈值 exit=", r2.returncode)
    check("放宽后仅告警（exit=0）", r2.returncode == 0, str(r2.returncode))
    check("放宽后输出含「未超上限」", "未超上限" in (r2.stdout + r2.stderr), (r2.stdout + r2.stderr)[-300:])

print("\n=== 结果 ===")
print("FAIL =", fail)
sys.exit(0 if fail == 0 else 1)
